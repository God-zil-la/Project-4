from io import BytesIO, StringIO
import csv

from docx import Document
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from xml.sax.saxutils import escape


SUPPORTED_FORMATS = {"pdf", "docx", "xlsx", "csv"}


def generate_document(title, content, file_format):
    """Generate a document in memory and return its bytes."""

    file_format = str(file_format).lower().strip()

    if file_format not in SUPPORTED_FORMATS:
        raise ValueError("Unsupported document format.")

    title = str(title).strip() or "Document"
    content = str(content)

    if file_format == "docx":
        document = Document()
        document.add_heading(title, 0)

        for paragraph in content.split("\n"):
            document.add_paragraph(paragraph)

        buffer = BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    if file_format == "pdf":
        buffer = BytesIO()
        pdf = SimpleDocTemplate(buffer, pagesize=A4)
        styles = getSampleStyleSheet()

        elements = [
            Paragraph(escape(title), styles["Title"]),
            Spacer(1, 16),
        ]

        for paragraph in content.split("\n"):
            elements.append(
                Paragraph(
                    escape(paragraph) or " ",
                    styles["Normal"],
                )
            )
            elements.append(Spacer(1, 8))

        pdf.build(elements)
        return buffer.getvalue()

    if file_format == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Document"

        sheet.append([title])

        for line in content.split("\n"):
            sheet.append([line])

        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    output = StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow([title])

    for line in content.split("\n"):
        writer.writerow([line])

    return output.getvalue().encode("utf-8-sig")
