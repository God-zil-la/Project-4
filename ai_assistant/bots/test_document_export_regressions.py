"""Local regression coverage for source-bound XLSX and readable exports."""

from io import BytesIO
import json
from unittest import TestCase
from unittest.mock import Mock, patch

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from .document_export_quality import clean_document_content, prepare_document_content
from .document_generator import generate_document
from .document_workbook_validation import bind_and_validate, extract_amounts


REQUEST = (
    "Skapa exakt tre ark 'M\u00e5nadsbudget', 'F\u00f6rs\u00e4ljningsprognos', "
    "'Resultatsammanst\u00e4llning'.\n"
    "hyra 15000 SEK, personal 25000, r\u00e5varor 10000, \u00f6vrigt 5000, "
    "SUM m\u00e5nad 55000, \u00e5rskostnad 660000 via formel; "
    "jan-dec 50000 varje m\u00e5nad, SUM \u00e5r 600000, resultat -60000"
)


def budget_spec():
    return {'sheets': [
        {'name': 'M\u00e5nadsbudget', 'rows': [
            ['Kostnad', 'SEK'],
            ['Hyra', {'source': 'amount_1'}],
            ['Personal', {'source': 'amount_2'}],
            ['R\u00e5varor', {'source': 'amount_3'}],
            ['\u00d6vrigt', {'source': 'amount_4'}],
            ['SUM m\u00e5nad', '=SUM(B2:B5)'],
            ['\u00c5rskostnad', '=B6*12'],
        ]},
        {'name': 'F\u00f6rs\u00e4ljningsprognos', 'rows': [
            ['M\u00e5nad', 'SEK'],
            *[[month, {'source': 'amount_7'}] for month in
              ('jan', 'feb', 'mar', 'apr', 'maj', 'jun', 'jul', 'aug', 'sep', 'okt', 'nov', 'dec')],
            ['SUM \u00e5r', '=SUM(B2:B13)'],
        ]},
        {'name': 'Resultatsammanst\u00e4llning', 'rows': [
            ['Post', 'SEK'],
            ['\u00c5rsint\u00e4kter', "='F\u00f6rs\u00e4ljningsprognos'!B14"],
            ['\u00c5rskostnad', "='M\u00e5nadsbudget'!B7"],
            ['Resultat', '=B2-B3'],
        ]},
    ]}



class SourceLabelRegressionTests(TestCase):
    def test_english_instruction_source_labels(self):
        message = (
            "Include Hyra 15 000 SEK, Personal 25 000 SEK, "
            "R\u00e5varor 10 000 SEK, and \u00d6vrigt 5 000 SEK."
        )
        amounts = extract_amounts(message)

        self.assertEqual(
            [item["label"] for item in amounts],
            ["Hyra", "Personal", "R\u00e5varor", "\u00d6vrigt"],
        )

        spec = {"sheets": [{"name": "Budget", "rows": [
            ["Hyra", {"source": "amount_1"}],
            ["Personal", {"source": "amount_2"}],
            ["R\u00e5varor", {"source": "amount_3"}],
            ["\u00d6vrigt", {"source": "amount_4"}],
            ["Total", "=SUM(B1:B4)"],
        ]}]}

        evaluator = bind_and_validate(spec, amounts)
        self.assertEqual(evaluator.cell("Budget", 5, 2), 55000)

    def test_swapped_source_binding_is_rejected(self):
        amounts = extract_amounts(
            "Include Hyra 15 000 SEK, Personal 25 000 SEK."
        )
        spec = {"sheets": [{"name": "Budget", "rows": [
            ["Hyra", {"source": "amount_2"}],
            ["Personal", {"source": "amount_1"}],
        ]}]}

        with self.assertRaisesRegex(
            ValueError, "Source label does not match its row"
        ):
            bind_and_validate(spec, amounts)



class ProductionPromptExtractionTests(TestCase):
    def test_full_english_production_amount_extraction(self):
        message = (
            "Create a downloadable Excel workbook with exactly three worksheets.\n"
            "Worksheet 1: M\\u00e5nadsbudget.\n"
            "Include Hyra 15 000 SEK, Personal 25 000 SEK, "
            "R\\u00e5varor 10 000 SEK, and \\u00d6vrigt 5 000 SEK.\n"
            "Calculate monthly total using an Excel formula and annual costs using a formula.\n"
            "Worksheet 2: F\\u00f6rs\\u00e4ljningsprognos.\n"
            "Include January through December with 50 000 SEK in sales for each month.\n"
            "Calculate annual sales using a formula.\n"
            "Worksheet 3: Resultatsammanst\\u00e4llning.\n"
            "Reference the totals from the other two worksheets using Excel formulas "
            "and calculate the annual result.\n"
            "Generate the actual downloadable .xlsx file, not a textual description."
        )

        amounts = extract_amounts(message)

        self.assertEqual(
            [item["value"] for item in amounts],
            ["15000", "25000", "10000", "5000", "50000"],
        )
        self.assertEqual(
            [item["label"] for item in amounts[:4]],
            ["Hyra", "Personal", "R\\u00e5varor", "\\u00d6vrigt"],
        )
        self.assertEqual(amounts[4]["count"], 12)
        self.assertFalse(amounts[4]["derived"])

    def test_worksheet_numbers_are_not_money(self):
        for number in (1, 2, 3):
            with self.subTest(number=number):
                self.assertEqual(
                    extract_amounts(f"Worksheet {number}: Budget."),
                    [],
                )

    def test_twelve_month_phrasing(self):
        cases = (
            "Include January through December with 50 000 SEK in sales for each month.",
            "Sales 50 000 SEK each month from January to December.",
            "Sales: 50 000 SEK per month for 12 months.",
            "January-December: 50 000 SEK each month.",
        )
        for message in cases:
            with self.subTest(message=message):
                amounts = extract_amounts(message)
                self.assertEqual(len(amounts), 1)
                self.assertEqual(amounts[0]["value"], "50000")
                self.assertEqual(amounts[0]["count"], 12)


class WorkbookRegressionTests(TestCase):
    def test_real_three_sheet_workbook_and_calculated_results(self):
        spec = budget_spec()
        evaluator = bind_and_validate(spec, extract_amounts(REQUEST))
        data = generate_document('Budget', json.dumps(spec), 'xlsx')
        workbook = load_workbook(BytesIO(data))
        self.assertEqual(workbook.sheetnames, [s['name'] for s in spec['sheets']])
        self.assertEqual([workbook.worksheets[0].cell(r, 2).value for r in range(2, 6)],
                         [15000, 25000, 10000, 5000])
        self.assertEqual([workbook.worksheets[1].cell(r, 2).value for r in range(2, 14)], [50000] * 12)
        for sheet, row, expected in ((0, 6, 55000), (0, 7, 660000), (1, 14, 600000), (2, 4, -60000)):
            self.assertEqual(evaluator.cell(workbook.sheetnames[sheet], row, 2), expected)
            self.assertEqual(workbook.worksheets[sheet].cell(row, 2).data_type, 'f')
        self.assertEqual(workbook.worksheets[2]['B3'].value, "='M\u00e5nadsbudget'!B7")


    def test_original_english_prompt_generates_valid_xlsx(self):
        message = (
            "Worksheet 1: M\u00e5nadsbudget.\n"
            "Include Hyra 15 000 SEK, Personal 25 000 SEK, "
            "R\u00e5varor 10 000 SEK, and \u00d6vrigt 5 000 SEK.\n"
            "Worksheet 2: F\u00f6rs\u00e4ljningsprognos.\n"
            "Include January through December with 50 000 SEK in sales for each month.\n"
            "Worksheet 3: Resultatsammanst\u00e4llning."
        )
        amounts = extract_amounts(message)
        self.assertEqual(len(amounts), 5)
        self.assertEqual(amounts[4]["count"], 12)

        spec = budget_spec()
        for row in spec["sheets"][1]["rows"]:
            if isinstance(row[1], dict) and row[1].get("source") == "amount_7":
                row[1]["source"] = "amount_5"

        evaluator = bind_and_validate(spec, amounts)
        data = generate_document("Budget", json.dumps(spec), "xlsx")
        workbook = load_workbook(BytesIO(data))

        self.assertEqual(
            workbook.sheetnames,
            [sheet["name"] for sheet in spec["sheets"]],
        )
        self.assertEqual(
            [workbook.worksheets[0].cell(r, 2).value for r in range(2, 6)],
            [15000, 25000, 10000, 5000],
        )
        self.assertEqual(
            [workbook.worksheets[1].cell(r, 2).value for r in range(2, 14)],
            [50000] * 12,
        )
        for index, row, expected in (
            (0, 6, 55000),
            (0, 7, 660000),
            (1, 14, 600000),
            (2, 4, -60000),
        ):
            sheet = workbook.worksheets[index]
            self.assertEqual(sheet.cell(row, 2).data_type, "f")
            self.assertEqual(evaluator.cell(sheet.title, row, 2), expected)

    def test_generic_labels_decimal_negative_and_small_currency(self):
        amounts = extract_amounts('Hosting: 1 250,50 SEK; Refund: -50 SEK; Seats: 3')
        spec = {'sheets': [{'name': 'Operations', 'rows': [
            ['Hosting', {'source': 'amount_1'}], ['Refund', {'source': 'amount_2'}],
            ['Seats', {'source': 'amount_3'}], ['Net', '=SUM(B1:B2)'],
        ]}]}
        evaluator = bind_and_validate(spec, amounts)
        self.assertEqual(str(evaluator.cell('Operations', 4, 2)), '1200.5')

    def test_altered_raw_amount_and_swapped_source_rejected(self):
        for value in (14000, 15000, {'source': 'amount_2'}, {'source': 'missing'}):
            with self.subTest(value=value), self.assertRaises(ValueError):
                spec = budget_spec()
                spec['sheets'][0]['rows'][1][1] = value
                bind_and_validate(spec, extract_amounts(REQUEST))

    def test_missing_month_and_wrong_cross_sheet_reference_rejected(self):
        for change in ('month', 'duplicate_month', 'reference', 'hardcoded', 'wrong_result'):
            with self.subTest(change=change), self.assertRaises(ValueError):
                spec = budget_spec()
                if change == 'month':
                    spec['sheets'][1]['rows'][12][1] = 0
                elif change == 'duplicate_month':
                    spec['sheets'][1]['rows'][12][0] = 'jan'
                elif change == 'reference':
                    spec['sheets'][2]['rows'][2][1] = "='M\u00e5nadsbudget'!B6"
                elif change == 'hardcoded':
                    spec['sheets'][0]['rows'][6][1] = {'source': 'amount_6'}
                else:
                    spec['sheets'][2]['rows'][3][1] = '=B3-B2'
                bind_and_validate(spec, extract_amounts(REQUEST))

    def test_monthly_cost_is_not_a_twelve_month_series(self):
        amounts = extract_amounts('Hosting: 1500 SEK per month')
        self.assertEqual(amounts[0]['count'], 1)

    def test_trailing_result_label_and_sentence_ending_amount(self):
        amounts = extract_amounts('Hyra 15000.\nKostnader och -60000 resultat')
        self.assertEqual(amounts[0]['value'], '15000')
        self.assertEqual(amounts[1]['label'], 'resultat')
        self.assertTrue(amounts[1]['derived'])

    def test_external_unsafe_missing_and_circular_formulas_rejected(self):
        for formula in ('=HYPERLINK("https://example.com")', "='[external.xlsx]Sheet'!A1",
                        "=cmd|' /C calc'!A0", '=WEBSERVICE("x")', '=B1', '=Z250',
                        '=1/0', '=SUM(B1:B999999)', '=INDIRECT("B2")', '=1+', '=NaN'):
            with self.subTest(formula=formula), self.assertRaises(ValueError):
                generate_document('Test', json.dumps({'sheets': [{'name': 'Data', 'rows': [['Value', formula]]}]}), 'xlsx')

    def test_malformed_and_limits_rejected(self):
        for spec in ({}, {'sheets': []}, {'sheets': [{'name': 'Data', 'rows': 'bad'}]},
                     {'sheets': [{'name': 'Data', 'rows': [[float('nan')]]}]},
                     {'sheets': [{'name': 'Data', 'rows': [[1] * 31]}]},
                     {'sheets': [{'name': 'Data', 'rows': [[1]]}] * 2}):
            with self.subTest(spec=spec), self.assertRaises(ValueError):
                generate_document('Test', json.dumps(spec), 'xlsx')

    def test_generic_no_money_workbook(self):
        spec = {'sheets': [{'name': 'Inventory', 'rows': [
            ['Item', 'Count'], ['Pencils', 3], ['Pens', 4], ['Total', '=SUM(B2:B3)']]}]}
        self.assertEqual(bind_and_validate(spec).cell('Inventory', 4, 2), 7)


class ProviderWorkbookTests(TestCase):
    @staticmethod
    def response(spec):
        return {'choices': [{'finish_reason': 'stop', 'message': {'content': json.dumps(spec)}}]}

    @patch('ai_assistant.bots.document_workbook.provider_call')
    def test_provider_to_actual_xlsx(self, provider):
        from .document_workbook import generate_workbook_spec
        provider.return_value = self.response(budget_spec())
        content = generate_workbook_spec(REQUEST, [])
        self.assertEqual(len(load_workbook(BytesIO(generate_document('Budget', content, 'xlsx'))).sheetnames), 3)
        self.assertEqual(provider.call_count, 1)


    @patch('ai_assistant.bots.document_workbook.provider_call')
    def test_original_english_prompt_provider_to_xlsx(self, provider):
        from .document_workbook import generate_workbook_spec

        message = (
            "Create a downloadable Excel workbook with exactly three worksheets. "
            "Worksheet 1: M\u00e5nadsbudget. "
            "Include Hyra 15 000 SEK, Personal 25 000 SEK, "
            "R\u00e5varor 10 000 SEK, and \u00d6vrigt 5 000 SEK. "
            "Calculate monthly total using an Excel formula and annual costs using a formula. "
            "Worksheet 2: F\u00f6rs\u00e4ljningsprognos. "
            "Include January through December with 50 000 SEK in sales for each month. "
            "Calculate annual sales using a formula. "
            "Worksheet 3: Resultatsammanst\u00e4llning. "
            "Reference the totals from the other two worksheets using Excel formulas "
            "and calculate the annual result. Generate the actual downloadable .xlsx file, "
            "not a textual description."
        )

        spec = budget_spec()
        for row in spec["sheets"][1]["rows"]:
            if isinstance(row[1], dict) and row[1].get("source") == "amount_7":
                row[1]["source"] = "amount_5"

        provider.return_value = self.response(spec)
        content = generate_workbook_spec(message, [])
        self.assertEqual(provider.call_count, 1)

        call = provider.call_args
        system_prompt = call.kwargs["messages"][0]["content"]
        self.assertIn('"count": 12', system_prompt)
        self.assertIn('"id": "amount_5"', system_prompt)

        workbook = load_workbook(
            BytesIO(generate_document("Budget", content, "xlsx"))
        )
        self.assertEqual(
            workbook.sheetnames,
            ["M\u00e5nadsbudget", "F\u00f6rs\u00e4ljningsprognos",
             "Resultatsammanst\u00e4llning"],
        )
        self.assertEqual(
            [workbook.worksheets[1].cell(r, 2).value for r in range(2, 14)],
            [50000] * 12,
        )
        self.assertEqual(workbook.worksheets[2]["B4"].value, "=B2-B3")
        self.assertEqual(workbook.worksheets[2]["B4"].data_type, "f")

    @patch('ai_assistant.bots.document_workbook.provider_call')
    def test_retry_and_fail_closed(self, provider):
        from .document_workbook import generate_workbook_spec
        altered = budget_spec()
        altered['sheets'][0]['rows'][1][1] = 14000
        provider.side_effect = [self.response(altered), self.response(budget_spec())]
        generate_workbook_spec(REQUEST, [])
        self.assertEqual(provider.call_count, 2)
        provider.side_effect = [self.response(altered), self.response(altered)]
        with self.assertRaises(ValueError):
            generate_workbook_spec(REQUEST, [])

    @patch('ai_assistant.bots.document_workbook.provider_call')
    def test_followup_uses_user_inputs_not_assistant_amounts(self, provider):
        from .document_workbook import generate_workbook_spec
        provider.return_value = self.response(budget_spec())
        generate_workbook_spec('Export as XLSX', [
            {'role': 'user', 'content': REQUEST}, {'role': 'assistant', 'content': 'Hyra 99999 SEK'}])

    @patch('ai_assistant.bots.document_workbook.provider_call')
    def test_wrong_sheet_names_and_incomplete_json_fail_closed(self, provider):
        from .document_workbook import generate_workbook_spec
        spec = budget_spec()
        spec['sheets'][2]['name'] = 'Wrong name'
        for response in (self.response(spec), {'choices': []},
                         {'choices': [{'finish_reason': 'length', 'message': {'content': '{}'}}]}):
            with self.subTest(response=response), self.assertRaises(ValueError):
                provider.return_value = response
                generate_workbook_spec(REQUEST, [])


class ExportBoundaryTests(TestCase):
    @patch('ai_assistant.bots.document_pipeline.upload_document')
    def test_wrong_owner_or_bot_denied_before_generation(self, upload):
        from .document_pipeline import create_chat_document
        user = Mock(pk=1, is_authenticated=True)
        for owner, conversation_user, conversation_bot in ((2, 1, 3), (1, 2, 3), (1, 1, 4)):
            with self.subTest(owner=owner, conversation_user=conversation_user, conversation_bot=conversation_bot):
                with self.assertRaises(PermissionError):
                    create_chat_document(user=user, bot=Mock(pk=3, owner_id=owner),
                                         conversation=Mock(user_id=conversation_user, bot_id=conversation_bot),
                                         title='Budget', content='invalid JSON', file_format='xlsx')
        upload.assert_not_called()

    @patch('ai_assistant.bots.document_pipeline.upload_document')
    def test_invalid_workbook_never_uploaded(self, upload):
        from .document_pipeline import create_chat_document
        with self.assertRaises(ValueError):
            create_chat_document(user=Mock(pk=1, is_authenticated=True),
                                 title='Budget', content='{"sheets":[]}', file_format='xlsx')
        upload.assert_not_called()


class ReadableExportTests(TestCase):
    trailer = 'Nu genererar jag PDF-filen med ovanst\u00e5ende inneh\u00e5ll. V\u00e4nligen v\u00e4nta en stund.'
    content = ('# Budget\n## Ekonomi\n**\u00c5rskostnad**\n'
               r'\(55 000 \times 12 = 660 000 \, \text{SEK}\)' '\n'
               '- Hyra\n1. Kontrollera resultat\n')

    def test_exact_swedish_trailer_and_factual_text(self):
        self.assertEqual(clean_document_content('Fakta\n' + self.trailer, strict=True), 'Fakta')
        for text in ('Jag skapar budgetar varje m\u00e5nad.', 'H\u00e4mta fakturor fr\u00e5n arkivet.',
                     'Nu genererar jag PDF-filen som exempel i utbildningen.'):
            self.assertEqual(clean_document_content('Fakta\n' + text, strict=True), 'Fakta\n' + text)

    def test_actual_docx_text_title_styles_and_math(self):
        data = generate_document('Budget', self.content + self.trailer, 'docx')
        document = Document(BytesIO(data))
        text = '\n'.join(p.text for p in document.paragraphs)
        self.assertEqual(text.count('Budget'), 1)
        self.assertIn('55 000 \u00d7 12 = 660 000 SEK', text)
        self.assertNotIn('genererar', text)
        self.assertNotIn('\\', text)
        self.assertIn('Heading 2', [p.style.name for p in document.paragraphs])
        self.assertIn('List Bullet', [p.style.name for p in document.paragraphs])
        self.assertIn('List Number', [p.style.name for p in document.paragraphs])
        self.assertTrue(any(r.bold and r.text == '\u00c5rskostnad' for p in document.paragraphs for r in p.runs))

    def test_actual_pdf_text_title_and_math(self):
        data = generate_document('Budget', self.content + self.trailer, 'pdf')
        text = '\n'.join(page.extract_text() for page in PdfReader(BytesIO(data)).pages)
        self.assertEqual(text.count('Budget'), 1)
        self.assertIn('55 000 \u00d7 12 = 660 000 SEK', text)
        self.assertIn('\u00c5rskostnad', text)
        self.assertNotIn('genererar', text)
        self.assertNotIn('\\', text)

    def test_fraction_and_unsupported_math(self):
        self.assertEqual(prepare_document_content('Test', r'\[\frac{10}{2} = 5\]'), '(10)/(2) = 5')
        with self.assertRaises(ValueError):
            prepare_document_content('Test', r'\(\unknown{10}\)')

    def test_csv_output_and_injection_guard_unchanged(self):
        data = generate_document('Data', 'produkt,pris_sek,antal,totalt_sek\nKaffe,25,2,50', 'csv')
        self.assertEqual(data.decode('utf-8-sig'), 'produkt,pris_sek,antal,totalt_sek\r\nKaffe,25,2,50\r\n')
        data = generate_document('Data', 'name,value,note\n=cmd,1,test', 'csv')
        self.assertIn("'=cmd", data.decode('utf-8-sig'))
